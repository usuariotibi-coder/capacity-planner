"""
Finance actuals import.

Parses a cumulative "WIP" Finance report (an .xlsx with columns Date, one column
per department, Labor Total, Job) and writes ProjectDepartmentWeeklyActual rows —
real historical hours per project/department/week, with no per-employee breakdown.

Key rules (confirmed with the business owner):
  - The report is cumulative: each report date's value for a given Job/department
    is a running total, not a weekly figure. The weekly figure is the diff against
    the previous report we have on file for that same Job/department.
  - A report dated D reflects data as of the END of the week before D's week (the
    report is generated and shared a week after the fact), so its "target week" is
    ISO-week(D) - 1.
  - The very first time we see a Job/department combination, there is nothing to
    diff against — that cumulative number could represent months of pre-existing
    work of unknown duration. We only seed a baseline in that case; no weekly
    actual is written for it.
  - If the gap between two consecutive reports for a Job/department spans more than
    one target week (a report was skipped), the diff is split evenly across the
    missing weeks.
  - Negative diffs (Finance corrections) are kept as-is, not clamped to zero.
  - Job codes collapse to a "project key": trailing lettered suffixes (e.g. "-A",
    "-YB") are dropped, but a "-CO##" change-order suffix is kept as part of the
    key (e.g. "3268-E" and "3268-A" both collapse to "3268"; "3268-CO01-A" would
    collapse to "3268-CO01"). Each project key maps 1:1 to a Project.project_number,
    auto-creating the Project if it doesn't exist yet.
  - Any existing Assignment rows covering a (project, department, week) that just
    received a fresh actual are deleted, so Finance's number is the sole source for
    that week (no double-counting when the two are summed for "Used" hours).
"""
from __future__ import annotations

import re
import datetime
from collections import defaultdict

import openpyxl
from django.db import transaction
from django.db.models import Q

from .models import (
    Project, ProjectDepartmentWeeklyActual, FinanceJobCumulative,
    FinanceImportLog, Assignment, Department,
)

# Excel header -> our Department code. Columns sharing a department are summed.
COLUMN_DEPARTMENT_MAP = {
    'PM': Department.PM,
    'Mechanical Engineering': Department.MED,
    'Hardware Design': Department.HD,
    'Machining': Department.MFG,
    'Assembly': Department.BUILD,
    'Assembly Install': Department.BUILD,
    'Programming': Department.PRG,
    'Programming Install': Department.PRG,
    'Purchasing': Department.PURCHASING,
    'Shipping': Department.PURCHASING,
}
IGNORED_COLUMNS = {'Labor Total'}

JOB_KEY_RE = re.compile(r'^(\d{3,6}(?:-CO\d+)?)')


def project_key_from_job(job_code: str) -> str | None:
    match = JOB_KEY_RE.match((job_code or '').strip())
    return match.group(1) if match else None


def _iso_monday(date_obj: datetime.date) -> datetime.date:
    """Monday of the ISO week containing date_obj."""
    return date_obj - datetime.timedelta(days=date_obj.weekday())


def _target_week(report_date: datetime.date) -> datetime.date:
    """The week a report date's cumulative figure actually reflects: last week."""
    return _iso_monday(report_date) - datetime.timedelta(weeks=1)


def _weeks_between(start_week: datetime.date, end_week: datetime.date) -> list[datetime.date]:
    """Mondays from start_week+1week through end_week, inclusive."""
    weeks = []
    w = start_week + datetime.timedelta(weeks=1)
    while w <= end_week:
        weeks.append(w)
        w += datetime.timedelta(weeks=1)
    return weeks


def parse_workbook(file_obj):
    """
    Returns {report_date: {job_code: {department: hours}}}, sorted-ready (caller
    sorts). Sums columns that map to the same department (e.g. Assembly + Assembly
    Install), and skips ignored columns (Labor Total).
    """
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.worksheets[0]

    headers = {}
    for col in range(1, ws.max_column + 1):
        headers[col] = ws.cell(row=1, column=col).value

    date_col = next((c for c, h in headers.items() if h == 'Date'), None)
    job_col = next((c for c, h in headers.items() if h == 'Job'), None)
    if not date_col or not job_col:
        raise ValueError("Expected 'Date' and 'Job' columns not found in the uploaded file.")

    dept_cols = {
        c: COLUMN_DEPARTMENT_MAP[h]
        for c, h in headers.items()
        if h in COLUMN_DEPARTMENT_MAP
    }

    data: dict[datetime.date, dict[str, dict[str, float]]] = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))

    for row in range(2, ws.max_row + 1):
        raw_date = ws.cell(row=row, column=date_col).value
        job_code = ws.cell(row=row, column=job_col).value
        if raw_date is None or job_code is None:
            continue
        report_date = raw_date.date() if isinstance(raw_date, datetime.datetime) else raw_date
        job_code = str(job_code).strip()
        if not job_code:
            continue

        for col, dept in dept_cols.items():
            value = ws.cell(row=row, column=col).value
            if value:
                data[report_date][job_code][dept] += float(value)

    return data


@transaction.atomic
def run_finance_import(file_obj, uploaded_by=None, file_name='') -> FinanceImportLog:
    parsed = parse_workbook(file_obj)
    report_dates = sorted(parsed.keys())

    # (project_key, department, week_start_date) -> accumulated hours this run
    actuals_accum: dict[tuple, float] = defaultdict(float)
    weeks_touched_by_project_dept: dict[tuple, set] = defaultdict(set)
    seeded_only_jobs = set()
    warnings = []

    # Preload existing cumulative baselines keyed by (job_code, department)
    baselines = {
        (row.job_code, row.department): row
        for row in FinanceJobCumulative.objects.all()
    }

    for report_date in report_dates:
        jobs_today = parsed[report_date]
        for job_code, dept_hours in jobs_today.items():
            project_key = project_key_from_job(job_code)
            if not project_key:
                warnings.append(f"Could not parse a project number from Job '{job_code}'; skipped.")
                continue

            for department, cumulative_value in dept_hours.items():
                baseline_key = (job_code, department)
                existing = baselines.get(baseline_key)

                if existing is None:
                    FinanceJobCumulative.objects.update_or_create(
                        job_code=job_code, department=department,
                        defaults={'cumulative_hours': cumulative_value, 'last_report_date': report_date},
                    )
                    baselines[baseline_key] = FinanceJobCumulative(
                        job_code=job_code, department=department,
                        cumulative_hours=cumulative_value, last_report_date=report_date,
                    )
                    seeded_only_jobs.add(f"{job_code} ({department})")
                    continue

                if report_date <= existing.last_report_date:
                    # Already-processed or out-of-order report date; skip safely.
                    continue

                diff = cumulative_value - existing.cumulative_hours
                start_week = _target_week(existing.last_report_date)
                end_week = _target_week(report_date)
                weeks = _weeks_between(start_week, end_week)

                if weeks and diff != 0:
                    per_week = diff / len(weeks)
                    for week in weeks:
                        key = (project_key, department, week)
                        actuals_accum[key] += per_week
                        weeks_touched_by_project_dept[(project_key, department)].add(week)

                existing.cumulative_hours = cumulative_value
                existing.last_report_date = report_date
                FinanceJobCumulative.objects.filter(job_code=job_code, department=department).update(
                    cumulative_hours=cumulative_value, last_report_date=report_date,
                )

    if not actuals_accum:
        log = FinanceImportLog.objects.create(
            uploaded_by=uploaded_by,
            file_name=file_name,
            report_dates_processed=[d.isoformat() for d in report_dates],
            new_projects_created=[],
            weekly_actuals_written=0,
            job_codes_seeded_only=sorted(seeded_only_jobs),
            warnings=warnings,
        )
        return log

    # Resolve/auto-create Project per project_key
    project_keys = sorted(set(k[0] for k in actuals_accum.keys()))
    projects_by_key = {
        p.project_number: p
        for p in Project.objects.filter(project_number__in=project_keys)
    }
    new_projects_created = []
    for key in project_keys:
        if key in projects_by_key:
            continue
        weeks_for_key = sorted({
            week for (pk, _dept, week) in actuals_accum.keys() if pk == key
        })
        start_date = weeks_for_key[0]
        end_date = weeks_for_key[-1] + datetime.timedelta(days=6)
        number_of_weeks = max(1, ((weeks_for_key[-1] - start_date).days // 7) + 1)
        project = Project.objects.create(
            project_number=key,
            name=key,
            client='',
            start_date=start_date,
            end_date=end_date,
            facility='MX',
            number_of_weeks=number_of_weeks,
        )
        projects_by_key[key] = project
        new_projects_created.append(key)

    # Write ProjectDepartmentWeeklyActual rows
    log = FinanceImportLog.objects.create(
        uploaded_by=uploaded_by,
        file_name=file_name,
        report_dates_processed=[d.isoformat() for d in report_dates],
        new_projects_created=new_projects_created,
        weekly_actuals_written=0,
        job_codes_seeded_only=sorted(seeded_only_jobs),
        warnings=warnings,
    )

    written = 0
    for (project_key, department, week), hours in actuals_accum.items():
        project = projects_by_key[project_key]
        ProjectDepartmentWeeklyActual.objects.update_or_create(
            project=project, department=department, week_start_date=week,
            defaults={'hours': hours, 'source_import': log},
        )
        written += 1

    # Finance replaces manual data for any (project, department, week) it just
    # supplied an actual for — remove the now-redundant Assignment rows so the two
    # sources aren't summed together.
    deleted_assignments = 0
    for (project_key, department), weeks in weeks_touched_by_project_dept.items():
        project = projects_by_key[project_key]
        # Effective department match: an explicit override, or the employee's own
        # department when no override is set.
        qs = Assignment.objects.filter(
            project=project,
            week_start_date__in=weeks,
        ).filter(
            Q(department_override=department)
            | (Q(department_override__isnull=True) & Q(employee__department=department))
        )
        count, _ = qs.delete()
        deleted_assignments += count

    log.weekly_actuals_written = written
    if deleted_assignments:
        log.warnings = warnings + [f"Removed {deleted_assignments} superseded manual Assignment rows."]
    log.save(update_fields=['weekly_actuals_written', 'warnings'])

    return log
